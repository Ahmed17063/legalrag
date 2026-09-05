from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_core.documents import Document as LCDocument
from docx import Document  
from typing import List

def process_docx(file_path: str) -> List[LCDocument]:
    """Extract structured content from DOCX files"""
    doc = Document(file_path)
    chunks = []
    current_heading = ""
    
    # Extract metadata from core properties
    metadata = {
        "source": file_path,
        "title": doc.core_properties.title,
        "author": doc.core_properties.author,
        "created": str(doc.core_properties.created),
        "modified": str(doc.core_properties.modified)
    }
    
    # Process document structure
    for para in doc.paragraphs:
        print("Paragraphe : ---------------")
        print(para.text)
        text = para.text.strip()
        if not text:
            continue
            
        # Detect heading styles
        if para.style.name.startswith('Heading'):
            current_heading = text
            continue
            
        # Create chunk with hierarchical context
        chunk_text = f"{current_heading}: {text}" if current_heading else text
        chunks.append(LCDocument(
            page_content=chunk_text,
            metadata=metadata.copy()
        ))
    
    # Process tables
    for table in doc.tables:
        table_text = "\n".join([" | ".join(cell.text for cell in row.cells) 
                               for row in table.rows])
        if table_text:
            chunks.append(LCDocument(
                page_content=f"{current_heading} [TABLE]: {table_text}",
                metadata=metadata.copy()
            ))
    
    # Smart chunking
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=50,
        separators=["\n\n", "\n", ". ", " ", ""]
    )
    
    return splitter.split_documents(chunks)